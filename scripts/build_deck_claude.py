#!/usr/bin/env python3
"""Build a PPTX lecture deck from a YAML content spec (Claude-built workflow).

    python3 scripts/build_deck_claude.py <spec.yaml> [--check] [--render]

Pipeline
  1. Validate the spec: every `verify` quote must occur in the syllabus text view, table weights must sum to
     `check_sum`, and slide minutes + Q&A must fit the configured in-class overhead from config/course.yaml.
  2. Build: open the course template (scripts/build_template_claude.py) for its master and layouts, remove
     its specimen slides, and create native, editable slides. Titles and plain bullet slides use the
     layout placeholders and inherit the master's style; content geometry is read from the template's body
     placeholder. Footer and slide number come from the master. Images are copied from named shapes of
     `image_source`, so provenance stays explicit and no image files are duplicated in the repository.
     Every generated text/table/image shape is named `yaml:<field path>` and each slide carries its spec id
     as the slide name, so edits made in PowerPoint can be mapped back to the spec.
  3. Estimate text overflow per text box (heuristic; warnings only).
  4. --render (macOS + Microsoft PowerPoint + pdftoppm): export to PDF through PowerPoint, rasterise, and
     write a contact sheet to the spec's `preview` path plus full-size PNGs for visual review. Each deck
     reuses one render folder (see --render-dir); its previous slide images and copies are replaced, so
     renders do not accumulate. Rendering is the real layout check; the overflow heuristic only warns.

--check validates the spec and exits without writing anything. Requires python-pptx, PyYAML and Pillow.
The template file and the frozen syllabus are only read.
Generated PPTX files go to output/deck_candidates/<name>.candidate.pptx. The
spec's output path supplies a name only. Existing candidates and working course/
decks are never overwritten; a reviewed candidate needs explicit promotion.
"""
import argparse, copy, io, math, re, subprocess, sys, tempfile, unicodedata
from pathlib import Path
from deck_paths import candidate_output, validate_candidate_output
import omml_claude as OMML

import yaml
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent

# Style follows the instructor's existing decks: white background, dark-red Calibri Light titles.
RED = RGBColor(0xC0, 0x00, 0x00)
INK = RGBColor(0x26, 0x26, 0x26)
MUTED = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
TINT = RGBColor(0xFB, 0xE9, 0xE9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_FONT, TITLE_FONT = 'Calibri', 'Calibri Light'
W, H = 13.333, 7.5
# Content frame (in); replaced in build() by the template's master body placeholder geometry.
MX, TOP, BOTTOM = 0.244, 0.845, 7.185
BODY, SMALL = 24, 16           # master body level-1 size; small text

warnings = []


# ---------------------------------------------------------------- validation

def norm(s):
    s = unicodedata.normalize('NFKC', s)
    s = re.sub(r'[‐-―−]', '-', s)
    s = s.replace('‘', "'").replace('’', "'").replace('“', '"').replace('”', '"')
    return re.sub(r'\s+', ' ', s).strip().lower()


def pct(s):
    return float(str(s).rstrip('%'))


def validate(spec):
    errors = []
    d = spec['deck']
    text = norm((ROOT / d['syllabus_text']).read_text(encoding='utf-8')) if d.get('syllabus_text') else ''
    for s in spec['slides']:
        for q in s.get('verify', []):
            if not text:
                errors.append(f"{s['id']}: verify quote given but the deck has no syllabus_text")
                continue
            if norm(q) not in text:
                errors.append(f"{s['id']}: verify quote not in {d['syllabus_text']}: {q!r}")
        t = s.get('table')
        if t and 'check_sum' in t:
            total = sum(pct(r[-1]) for r in t['rows'])
            if abs(total - t['check_sum']) > 1e-9 or pct(t['total'][-1]) != t['check_sum']:
                errors.append(f"{s['id']}: table weights sum to {total}, expected {t['check_sum']}")
    proj = next((s for s in spec['slides'] if s['kind'] == 'project'), None)
    asmt = next((s for s in spec['slides'] if s.get('id') == 'assessment'), None)
    if proj and asmt:
        row = next(r for r in asmt['table']['rows'] if 'project' in r[0].lower())
        if pct(row[-1]) != proj['table']['check_sum']:
            errors.append('project milestone total differs from the assessment-table project weight')
    budget = time_budget(d)
    timed = [s for s in spec['slides'] if not s.get('appendix')]
    for s in timed:
        if not isinstance(s.get('minutes'), (int, float)):
            errors.append(f"{s['id']}: no `minutes` estimate (needed for the {budget}-minute budget)")
    used = sum(s.get('minutes') or 0 for s in timed)
    total = used + d.get('qa_minutes', 0)
    ok = d.get('time_override', {})                   # instructor-accepted overrun, recorded in the spec
    if total > budget and total <= ok.get('accepted_minutes', budget):
        warnings.append(f'timing: {total} min exceeds the {budget}-minute budget; accepted up to '
                        f'{ok["accepted_minutes"]} min ({ok.get("authority", "no authority given")})')
    elif total > budget:
        errors.append(f'timing: {used} slide min + {d.get("qa_minutes", 0)} Q&A > {budget} configured minutes')
    return errors, used, budget


def time_budget(d):
    """Minutes available: a lecture session's accepted teaching minutes (`session_id` in
    decisions/topic_decisions.yaml) or a configured in-class overhead (`overhead_id` in config/course.yaml)."""
    if d.get('session_id'):
        plan = yaml.safe_load((ROOT / 'decisions/topic_decisions.yaml').read_text())['lecture_plan']
        return next(e['teaching_minutes'] for e in plan if e['id'] == d['session_id'])
    course = yaml.safe_load((ROOT / 'config/course.yaml').read_text())
    return next(o['minutes'] for o in course['design']['in_class_overheads'] if o['id'] == d['overhead_id'])


# ---------------------------------------------------------------- text helpers

# Inline markup: **bold**, *italic*, [text](url), <https://bare.url>
TOKEN = re.compile(r'\*\*(.+?)\*\*|\[(.+?)\]\((.+?)\)|<(https?://[^>\s]+)>|\*(.+?)\*')


def markup(text):
    """Split inline markup into (text, bold, italic, link) segments."""
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            yield text[pos:m.start()], False, False, None
        if m.group(1):
            yield m.group(1), True, False, None
        elif m.group(2):
            yield m.group(2), False, False, m.group(3)
        elif m.group(4):
            yield m.group(4), False, False, m.group(4)
        else:
            yield m.group(5), False, True, None
        pos = m.end()
    if pos < len(text):
        yield text[pos:], False, False, None


def add_runs(p, text, size, color=INK, bold=False, font=BODY_FONT):
    """Inline markup; $...$ becomes text-math runs (italic letters, sub/superscript baselines)."""
    parts = text.split('$')
    if len(parts) > 1:
        for k, part in enumerate(parts):
            if k % 2:
                text_math(p, part, size, color, bold)
            elif part:
                add_runs(p, part, size, color, bold, font)
        return
    for seg, b, it, link in markup(text):
        r = _run(p, seg, size, RGBColor(0x05, 0x63, 0xC1) if link else color, bold or b, font)
        if it:
            r.font.italic = True
        if link:
            r.hyperlink.address = link


def text_math(p, latex, size, color=INK, bold=False):
    """Render a small LaTeX expression as ordinary runs (for table cells, where native equations are avoided)."""
    def emit(items, baseline=0, sz=size):
        for it in items:
            if it[0] == 't':
                r = _run(p, it[1], sz, color, bold, 'Cambria Math')
                r.font.italic = not it[2]
                if baseline:
                    r.font._rPr.set('baseline', str(baseline))
            elif it[0] == 'grp':
                emit(it[1], baseline, sz)
            elif it[0] == 'script':
                emit([it[1]], baseline, sz)
                if it[2]:
                    emit(it[2], -25000, sz)
                if it[3]:
                    emit(it[3], 30000, sz)
            elif it[0] == 'nary':
                emit([('t', '∑', True)], baseline, sz); emit(it[3], baseline, sz)
    emit(OMML.Parser(latex).seq())


def _run(p, text, size, color, bold, font):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.name = Pt(size), bold, font
    f.color.rgb = color
    return r


def plain(text):
    return ''.join(seg for seg, *_ in markup(text))


def set_bullet(p, level, size):
    pPr = p._p.get_or_add_pPr()
    hang = 0.25                                    # same hang and level step as the master body style
    pPr.set('marL', str(int(Inches(hang + 0.31 * level))))
    pPr.set('indent', str(-int(Inches(hang))))
    for tag in ('a:buNone', 'a:buChar', 'a:buFont', 'a:buClr'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = pPr.makeelement(qn('a:buClr'), {})
    srgb = clr.makeelement(qn('a:srgbClr'), {'val': 'C00000' if level == 0 else '7F7F7F'})
    clr.append(srgb)
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
    pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•' if level == 0 else '–'}))


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0           # text edge = shape edge, aligned with the title
    tf.margin_top = tf.margin_bottom = 0
    return tb


def flatten(items, level=0):
    for it in items:
        if isinstance(it, list):
            yield from flatten(it, level + 1)
        else:
            yield level, it


def bullets(slide, x, y, w, h, items, size=BODY, gap=8, name='Body', style='bullet', anchor=MSO_ANCHOR.TOP):
    tb = textbox(slide, x, y, w, h, anchor=anchor, name=name)
    tf = tb.text_frame
    lines = []
    for i, (lvl, text) in enumerate(flatten(items)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl                                 # kept for PowerPoint indenting and the pull step
        s = size if lvl == 0 else size - 4
        add_runs(p, text, s, INK if lvl == 0 else MUTED)
        if style == 'bullet':
            set_bullet(p, lvl, s)
        p.space_before = Pt(gap if i and lvl == 0 else 2)
        lines.append((plain(text), s, 0.25 + 0.31 * lvl, gap if lvl == 0 else 2))
    check_fit(slide, name, w, h, lines)
    return tb


def para_box(slide, x, y, w, h, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font=BODY_FONT, name='Text'):
    tb = textbox(slide, x, y, w, h, anchor, name)
    p = tb.text_frame.paragraphs[0]
    p.alignment = align
    add_runs(p, text, size, color, bold, font)
    check_fit(slide, name, w, h, [(plain(text), size, 0, 0)])
    return tb


def check_fit(slide, name, w, h, lines):
    """Rough overflow estimate: Calibri averages ~0.5 em per character."""
    need = 0.06
    for text, size, indent, gap in lines:
        per_line = max(1, int((w - indent) * 72 / (0.5 * size)))
        need += math.ceil(max(1, len(text)) / per_line) * size * 1.2 / 72 + gap / 72
    if need > h + 0.05:
        warnings.append(f'slide {slide._idx}: "{name}" may overflow (~{need:.2f}in needed, {h:.2f}in available)')


# ---------------------------------------------------------------- slide parts

def set_title(slide, text):
    """Plain text in the layout's title placeholder: position, font, size and colour come from the master."""
    t = slide.shapes.title
    t.name = 'yaml:title'
    t.text_frame.text = text


def footnote(slide, text, y=None):
    para_box(slide, MX, y or BOTTOM - 0.45, W - 2 * MX, 0.45, text, SMALL + 2, MUTED, name='yaml:footnote',
             anchor=MSO_ANCHOR.BOTTOM)


def picture(slide, blob, x, y, w=None, h=None, name='yaml:image'):
    pic = slide.shapes.add_picture(io.BytesIO(blob), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    pic.name = name
    return pic


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = 'deco:box'
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def table(slide, x, y, w, spec, size=18, row_h=0.5, name='yaml:table'):
    rows = [spec['header']] + spec['rows'] + ([spec['total']] if 'total' in spec else [])
    widths = spec['widths']
    scale = w / sum(widths)
    gt = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    gt.name = name
    tbl = gt.table
    tbl.first_row = True
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Inches(cw * scale)
    last = len(rows) - 1
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = RED if i == 0 else (TINT if 'total' in spec and i == last else
                                                        (WHITE if i % 2 else LIGHT))
            c.margin_left = c.margin_right = Inches(0.12)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ''
            p.alignment = (PP_ALIGN.RIGHT if j == len(row) - 1 and j > 0 and not spec.get('align_left')
                           else PP_ALIGN.LEFT)
            header, total = i == 0, 'total' in spec and i == last
            add_runs(p, str(val), size, WHITE if header else INK, header or total or j == 0)
    # Plain table style: no theme banding, so our fills are the only fills.
    tblPr = tbl._tbl.tblPr
    for attr in ('bandRow', 'firstRow'):
        tblPr.set(attr, '0')
    return gt


def notes(slide, s):
    srcs = eq_sources(s.get('body')) + eq_sources(s.get('after'))
    if srcs:
        s = dict(s, notes=(s.get('notes', '').rstrip() + '\n\nEquation sources (LaTeX): ' + ' | '.join(srcs) + '\n'))
    head = f"Target: {s.get('minutes', 0)} min{' (appendix)' if s.get('appendix') else ''}. Sources: {', '.join(map(str, s.get('sources', [])))}."
    if s.get('instructor_voice'):
        head += f" Instructor voice (not syllabus policy): {'; '.join(s['instructor_voice'])}."
    slide.notes_slide.notes_text_frame.text = s.get('notes', '').strip() + '\n\n' + head


# ---------------------------------------------------------------- slide kinds
# Shape names `yaml:<path>` mirror the spec field each shape renders; `deco:*` shapes are decoration.

def k_title(slide, s, img):
    # Full-height image panel on the right, cropped (not distorted) around a horizontal focus point.
    pw = s.get('image_width', 5.8)
    pic = picture(slide, img(s['image']), W - pw, 0, w=pw, h=H)
    from PIL import Image
    iw, ih = Image.open(io.BytesIO(img(s['image']))).size
    keep = (pw / H) * ih / iw                                  # fraction of image width that fits
    if keep < 1:
        left = min(max(s.get('image_focus', 0.5) - keep / 2, 0), 1 - keep)
        pic.crop_left, pic.crop_right = left, 1 - keep - left
    tw = W - pw - MX - 0.4
    t = slide.shapes.title                                     # Title Slide layout: style from the layout
    t.name = 'yaml:title'
    t.left, t.top, t.width, t.height = Inches(MX), Inches(2.0), Inches(tw), Inches(2.0)
    t.text_frame.text = s['title']
    line = slide.shapes.add_connector(1, Inches(MX), Inches(4.2), Inches(MX + 1.2), Inches(4.2))
    line.name = 'deco:accent'
    line.line.color.rgb, line.line.width = RED, Pt(2.5)
    sub = next(ph for ph in slide.placeholders if ph.placeholder_format.type == 4)   # subtitle
    sub.name = 'yaml:subtitle'
    sub.left, sub.top, sub.width, sub.height = Inches(MX), Inches(4.4), Inches(tw), Inches(0.9)
    sub.text_frame.text = s['subtitle']
    para_box(slide, MX, 5.35, tw, 0.5, s['byline'], 20, MUTED, name='yaml:byline')


def k_course_title(slide, s, img):
    # Course title slide (instructor's format, Introduction deck 2026-10-08): course name, then the deck topic
    # (bold), instructor and term as centred 36 pt lines; geometry and colours otherwise from the layout.
    d = s['_deck']
    t = slide.shapes.title
    t.name = 'yaml:title'
    t.left, t.top, t.width, t.height = Emu(1133575), Emu(865991), Emu(9924847), Emu(1573306)
    sub = next(ph for ph in slide.placeholders if ph.placeholder_format.type == 4)   # subtitle
    sub.name = 'yaml:subtitle'                                 # topic, instructor, term lines
    sub.left, sub.top, sub.width, sub.height = Emu(2946392), Emu(2922493), Emu(6299215), Emu(2277035)
    for shape, rows in ((t, [(s['title'], None, False)]),
                        (sub, [(s['topic'], 36, True), (d['author'], 36, False), (d['term'], 36, False)])):
        tf = shape.text_frame
        bp = tf._txBody.bodyPr
        for child in list(bp):
            bp.remove(child)
        bp.append(bp.makeelement(qn('a:normAutofit'), {}))
        for i, (text, size, bold) in enumerate(rows):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = PP_ALIGN.CENTER
            r = par.add_run()
            r.text = text
            if size:
                r.font.size = Pt(size)
            if bold:
                r.font.bold = True


def k_image_bullets(slide, s, img):
    """Optional: image_width (in), valign: middle (text centred on the image), bullet_style: none."""
    right = s.get('image_side') == 'right'
    iw = s.get('image_width', 3.1)
    ix = W - MX - iw if right else MX
    pic = picture(slide, img(s['image']), ix, TOP, w=iw)
    if pic.height.inches > BOTTOM - TOP:                       # keep tall images inside the content frame
        scale = (BOTTOM - TOP) / pic.height.inches
        pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
        iw = pic.width.inches
        pic.left = Inches(W - MX - iw if right else MX)
    if s.get('image_border', True):
        pic.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
    if s.get('caption'):
        para_box(slide, ix - 0.3, TOP + 0.1 + pic.height.inches, iw + 0.6, 0.75, s['caption'], SMALL,
                 MUTED, align=PP_ALIGN.CENTER, name='yaml:caption')
    bx = MX if right else MX + iw + 0.5
    bw = W - 2 * MX - iw - 0.5
    if s.get('valign') == 'middle':
        bullets(slide, bx, TOP, bw, pic.height.inches, s['bullets'], name='yaml:bullets',
                style=s.get('bullet_style', 'bullet'), anchor=MSO_ANCHOR.MIDDLE)
    else:
        bullets(slide, bx, TOP, bw, 4.9, s['bullets'], name='yaml:bullets', style=s.get('bullet_style', 'bullet'))
    if s.get('footnote'):
        footnote(slide, s['footnote'])


def k_bullets(slide, s, img):
    """Uses the Title and Content body placeholder: bullet style and geometry come from the master."""
    body = next(ph for ph in slide.placeholders if ph.placeholder_format.idx == 1)
    body.name = 'yaml:bullets'
    tf = body.text_frame
    for i, (lvl, text) in enumerate(flatten(s['bullets'])):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = lvl
        for seg, b, it, link in markup(text):               # markup only; no size/colour overrides
            r = p.add_run()
            r.text = seg
            if b:
                r.font.bold = True
            if it:
                r.font.italic = True
            if link:
                r.hyperlink.address = link
    if s.get('footnote'):
        footnote(slide, s['footnote'])


def k_two_column(slide, s, img):
    cw = (W - 2 * MX - 0.4) / 2
    for i, key in enumerate(('left', 'right')):
        col, x = s[key], MX + i * (cw + 0.4)
        box(slide, x, TOP, cw, 0.55, RED if i == 0 else INK)
        para_box(slide, x + 0.12, TOP, cw - 0.24, 0.55, col['heading'], 20, WHITE, True,
                 anchor=MSO_ANCHOR.MIDDLE, name=f'yaml:{key}.heading')
        bullets(slide, x, TOP + 0.75, cw, 4.6, col['bullets'], gap=12, name=f'yaml:{key}.bullets')
    if s.get('footnote'):
        footnote(slide, s['footnote'])


def k_flow(slide, s, img):
    para_box(slide, MX, TOP, W - 2 * MX, 0.5, s['lead'], 22, INK, name='yaml:lead')
    n = s['columns']
    gap = 0.45
    bw = (W - 2 * MX - gap * (n - 1)) / n
    bh = 2.0
    for i, (head, detail) in enumerate(s['boxes']):
        r, c = divmod(i, n)
        x, y = MX + c * (bw + gap), TOP + 0.8 + r * (bh + 0.45)
        first_third = r == 0 and c < 2
        box(slide, x, y, bw, bh, TINT if first_third else LIGHT)
        box(slide, x, y, bw, 0.08, RED if first_third else MUTED)
        para_box(slide, x + 0.12, y + 0.16, bw - 0.24, 0.7, f'{i + 1}. {head}', 20, INK, True,
                 anchor=MSO_ANCHOR.MIDDLE, name=f'yaml:boxes.{i}.0')
        para_box(slide, x + 0.12, y + 0.95, bw - 0.24, 0.95, detail, 18, MUTED, name=f'yaml:boxes.{i}.1')
        if c < n - 1:
            a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.08), Inches(y + bh / 2 - 0.15),
                                       Inches(gap - 0.16), Inches(0.3))
            a.name = 'deco:arrow'
            a.fill.solid(); a.fill.fore_color.rgb = RGBColor(0xBF, 0xBF, 0xBF); a.line.fill.background()
    if s.get('footnote'):
        footnote(slide, s['footnote'])


def k_table(slide, s, img):
    """Optional: table.size, table.row_h, bullets, footnote."""
    t = s['table']
    rows = len(t['rows']) + 1 + ('total' in t)
    row_h = t.get('row_h', 0.68)
    table(slide, MX, TOP, W - 2 * MX, t, size=t.get('size', 22), row_h=row_h)
    y = TOP + rows * row_h + 0.4
    if s.get('bullets'):
        bullets(slide, MX, y, W - 2 * MX, BOTTOM - y, s['bullets'], size=22, name='yaml:bullets',
                style=s.get('bullet_style', 'bullet'))
    if s.get('footnote'):
        footnote(slide, s['footnote'])


def k_project(slide, s, img):
    n = len(s['steps'])
    gap = -0.12
    cw = (W - 2 * MX - gap * (n - 1)) / n
    for i, step in enumerate(s['steps']):
        x = MX + i * (cw + gap)
        ch = slide.shapes.add_shape(MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON,
                                    Inches(x), Inches(TOP), Inches(cw), Inches(1.0))
        ch.name = f'yaml:steps.{i}'
        ch.fill.solid(); ch.fill.fore_color.rgb = RED if i % 2 == 0 else RGBColor(0x9E, 0x1B, 0x1B)
        ch.line.fill.background(); ch.shadow.inherit = False
        tf = ch.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        _run(p, step, 17, WHITE, True, BODY_FONT)
    tw = 5.2
    bullets(slide, MX, TOP + 1.4, W - 2 * MX - tw - 0.5, 4.5, s['bullets'], size=22, gap=12, name='yaml:bullets')
    table(slide, W - MX - tw, TOP + 1.4, tw, s['table'], size=20, row_h=0.6)


def k_callouts(slide, s, img):
    cw = 3.1
    for i, (num, label) in enumerate(s['callouts']):
        y = TOP + i * 2.15
        box(slide, W - MX - cw, y, cw, 1.95, TINT)
        para_box(slide, W - MX - cw, y + 0.05, cw, 1.15, num, 66, RED, True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE, name=f'yaml:callouts.{i}.0')
        para_box(slide, W - MX - cw, y + 1.25, cw, 0.5, label, 20, INK, align=PP_ALIGN.CENTER,
                 name=f'yaml:callouts.{i}.1')
    bullets(slide, MX, TOP, W - 2 * MX - cw - 0.5, 5.6, s['bullets'], gap=12, name='yaml:bullets')


def k_closing(slide, s, img):
    """Section Header layout (no rule/footer); title and subtitle centred."""
    t = slide.shapes.title
    t.name = 'yaml:title'
    t.left, t.top, t.width, t.height = Inches(MX), Inches(2.4), Inches(W - 2 * MX), Inches(1.4)
    t.text_frame.text = s['title']
    t.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    sub = next(ph for ph in slide.placeholders if ph.placeholder_format.idx == 1)
    sub.name = 'yaml:subtitle'
    sub.left, sub.top, sub.width, sub.height = Inches(MX), Inches(4.1), Inches(W - 2 * MX), Inches(0.7)
    sub.text_frame.text = s['subtitle']
    sub.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER


def rich_paragraphs(tf, items, size, eq_size, first=True):
    """Paragraphs for the `math` kind: strings are bullet paragraphs with native inline math ($...$) and markup;
    {eq: latex} is a centred native display equation; {text: str} is an unbulleted paragraph; {sub: [...]}
    holds level-1 bullets."""
    for item in items:
        if isinstance(item, dict) and 'sub' in item:
            first = rich_paragraphs(tf, [{'_lvl': 1, 'p': x} for x in item['sub']], size, eq_size, first)
            continue
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        if isinstance(item, dict) and 'eq' in item:
            new = OMML.math_paragraph(item['eq'], item.get('size', eq_size))
            p._p.addnext(new); p._p.getparent().remove(p._p)
            continue
        lvl = item.get('_lvl', 0) if isinstance(item, dict) else 0
        text = item if isinstance(item, str) else item.get('p', item.get('text'))
        sz = size if lvl == 0 else size - 4
        if text.startswith('$'):
            _run(p, '\u2060', sz, INK, False, BODY_FONT)   # word joiner: PowerPoint drops the bullet before an equation
        for k, part in enumerate(text.split('$')):
            if k % 2:
                p._p.append(OMML.inline_math(part, sz))
            elif part:
                add_runs(p, part, sz, INK if lvl == 0 else MUTED)
        if isinstance(item, dict) and 'text' in item:
            pPr = p._p.get_or_add_pPr(); pPr.set('marL', '0'); pPr.set('indent', '0')
            pPr.append(pPr.makeelement(qn('a:buNone'), {}))
        else:
            set_bullet(p, lvl, sz)
        p.space_before = Pt(8 if lvl == 0 else 2)
    return first


def eq_sources(items):
    """LaTeX of a slide's display equations and $inline$ math, in order and without repeats (for the notes)."""
    out = []
    for item in items or []:
        if isinstance(item, dict) and 'eq' in item:
            out.append(item['eq'])
        elif isinstance(item, dict) and 'sub' in item:
            out += eq_sources(item['sub'])
        elif isinstance(item, dict) and 'text' in item:
            out += item['text'].split('$')[1::2]
        elif isinstance(item, str):
            out += item.split('$')[1::2]
    return list(dict.fromkeys(out))


def copy_figure(slide, fig):
    """Copy the shapes of a figure drawn natively on another deck's slide (by slide name), keeping positions and
    remapping shape ids so connectors stay attached. Copies are named `native:fig:<original name>`."""
    slides = Presentation(ROOT / fig['deck']).slides
    src = (slides[fig['slide'] - 1] if isinstance(fig['slide'], int)              # old decks: PPTX position
           else next(s for s in slides if s._element.cSld.get('name') == fig['slide']))
    tree = slide.shapes._spTree
    used = {int(i) for i in tree.xpath('.//p:cNvPr/@id')}
    nxt = max(used | {1}) + 1
    els, idmap = [], {}
    for sh in src.shapes:
        if sh.name in fig.get('exclude', []):
            continue
        if sh.shape_type == 13:                         # picture: re-add the image (relationships cannot be copied)
            pic = slide.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
            pic.name = 'native:fig:' + sh.name
            for attr in ('crop_left', 'crop_right', 'crop_top', 'crop_bottom'):
                setattr(pic, attr, getattr(sh, attr))
            nxt = max(nxt, pic.shape_id + 1)
            continue
        el = copy.deepcopy(sh._element)
        if el.xpath('.//@r:embed | .//@r:id | .//@r:link'):
            raise ValueError(f'figure shape {sh.name!r} has relationships; copy it through PowerPoint')
        for c in el.xpath('.//p:cNvPr'):
            idmap[c.get('id')] = str(nxt)
            if c is el.xpath('.//p:cNvPr')[0]:
                c.set('name', 'native:fig:' + c.get('name'))
            c.set('id', str(nxt)); nxt += 1
        els.append(el)
    for el in els:
        for cx in el.xpath('.//a:stCxn | .//a:endCxn'):
            cx.set('id', idmap.get(cx.get('id'), cx.get('id')))
    if fig.get('group'):                                # one group placed at [x, y] (optionally scaled to [.., w, h])
        tree.append(group_shapes(els, fig['group'], nxt))
    else:
        for el in els:
            tree.append(el)
    if fig.get('relabel'):
        relabel(slide, fig['relabel'])
    if fig.get('variant') == 'terminal':
        terminal_variant(slide, fig)


def group_shapes(els, box, gid):
    """Wrap copied shape elements in a p:grpSp whose top-left is box[:2] (inches); box[2:] scales it if given."""
    boxes = []
    for el in els:                                      # visual bounds: 90/270-degree rotations swap the extents
        f = el.xpath('./p:spPr/a:xfrm | ./p:grpSpPr/a:xfrm')[0]
        x, y = int(f.find(qn('a:off')).get('x')), int(f.find(qn('a:off')).get('y'))
        w_, h_ = int(f.find(qn('a:ext')).get('cx')), int(f.find(qn('a:ext')).get('cy'))
        if round(int(f.get('rot', 0)) / 5400000) % 2:
            x, y, w_, h_ = x + (w_ - h_) // 2, y + (h_ - w_) // 2, h_, w_
        boxes.append((x, y, x + w_, y + h_))
    x0, y0 = min(b[0] for b in boxes), min(b[1] for b in boxes)
    cx, cy = max(b[2] for b in boxes) - x0, max(b[3] for b in boxes) - y0
    w, h = (Inches(box[2]), Inches(box[3])) if len(box) == 4 else (cx, cy)
    g = etree.fromstring(
        f'<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:nvGrpSpPr><p:cNvPr id="{gid}" name="native:fig:group"/>'
        f'<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="{Inches(box[0])}" y="{Inches(box[1])}"/>'
        f'<a:ext cx="{w}" cy="{h}"/><a:chOff x="{x0}" y="{y0}"/><a:chExt cx="{cx}" cy="{cy}"/></a:xfrm>'
        '</p:grpSpPr></p:grpSp>')
    for el in els:
        g.append(el)
    return g


def relabel(slide, mapping):
    """Replace the text of copied shapes (also inside groups) whose whole text equals a key. The new text may hold
    $...$ (typeset as text runs) and '\\n' line breaks; the first run's size, colour and bold are kept."""
    from pptx.text.text import TextFrame
    found = set()
    for txBody in slide.shapes._spTree.iter(qn('p:txBody')):
        tf = TextFrame(txBody, None)
        old = tf.text.strip()
        if old not in mapping:
            continue
        found.add(old)
        r0 = next((r for p in tf.paragraphs for r in p.runs), None)
        size = r0.font.size.pt if r0 is not None and r0.font.size else 18
        color = r0.font.color.rgb if r0 is not None and r0.font.color and r0.font.color.type else INK
        bold = bool(r0.font.bold) if r0 is not None else False
        paras = tf.paragraphs
        for p in paras[1:]:
            p._p.getparent().remove(p._p)
        p0 = paras[0]
        for r in list(p0._p.xpath('./a:r | ./a:br | ./a:fld')):
            p0._p.remove(r)
        for k, line in enumerate(mapping[old].split('\n')):
            p = p0 if k == 0 else tf.add_paragraph()
            if k:
                p.alignment = p0.alignment
            add_runs(p, line, size, color, bold)
    missing = set(mapping) - found
    if missing:
        raise ValueError(f'relabel: no copied shape has the text {sorted(missing)}')


def terminal_variant(slide, fig):
    """Replace the rescue branch (failed low search -> high) by an 'out of battery' terminal state."""
    by = {sh.name: sh for sh in slide.shapes}
    for name in fig['remove']:
        el = by['native:fig:' + name]._element
        el.getparent().remove(el)
    t = fig['terminal']
    node = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(t['x']), Inches(t['y']), Inches(t['w']), Inches(t['h']))
    node.name = 'native:fig:terminal'
    node.fill.solid(); node.fill.fore_color.rgb = TINT
    node.line.color.rgb = RED; node.line.width = Pt(2)
    node.shadow.inherit = False
    tf = node.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _run(p, t['label'], t.get('size', 16), RED, True, BODY_FONT)
    src = by['native:fig:' + t['from']]
    sx, sy = src.left + src.width // 2, src.top + src.height // 2
    ex, ey = node.left + node.width, node.top + node.height // 2
    arrow = slide.shapes.add_connector(1, sx, sy, ex, ey)          # straight connector
    arrow.name = 'native:fig:terminal_arrow'
    arrow.begin_connect(src, 0); arrow.end_connect(node, 3)
    arrow.line.color.rgb = RED; arrow.line.width = Pt(1.5)
    ln = arrow.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle'}))
    for name, (x, y) in fig.get('move', {}).items():
        sh = by['native:fig:' + name]; sh.left, sh.top = Inches(x), Inches(y)


def k_math(slide, s, img):
    """Text with native equations, optional table and optional copied figure. Boxes are [x, y, w(, h)] in inches."""
    size, eq_size = s.get('size', 24), s.get('eq_size', s.get('size', 24))
    if s.get('figure'):
        copy_figure(slide, s['figure'])
    bx = s.get('body_box', [MX, TOP, W - 2 * MX, BOTTOM - TOP])
    if s.get('body'):
        tb = textbox(slide, *bx[:3], bx[3] if len(bx) > 3 else BOTTOM - bx[1], name='yaml:body')
        rich_paragraphs(tb.text_frame, s['body'], size, eq_size)
    if s.get('table'):
        tx = s.get('table_box', [MX, TOP + 1.2, W - 2 * MX])
        t = s['table']
        table(slide, tx[0], tx[1], tx[2], t, size=t.get('size', 20), row_h=t.get('row_h', 0.55))
    if s.get('after'):
        ax = s['after_box']
        tb = textbox(slide, *ax[:3], ax[3] if len(ax) > 3 else BOTTOM - ax[1], name='yaml:after')
        rich_paragraphs(tb.text_frame, s['after'], size, eq_size)
    if s.get('footnote'):
        footnote(slide, s['footnote'])


def k_pptx_only(slide, s, img):
    """Placeholder for a slide authored and maintained only in the instructor's PPTX: a rebuild shows where it
    belongs instead of silently dropping it."""
    para_box(slide, MX, TOP + 1.5, W - 2 * MX, 1.5, 'Instructor slide maintained in the course PPTX; '
             'not reproduced by the builder.', 24, MUTED, align=PP_ALIGN.CENTER, name='deco:pptx_only')


def k_blank(slide, s, img):
    """Title only; the canvas stays empty for live work."""


def k_image_grid(slide, s, img):
    """2×2 (or n-column) cells: image or drawn `comparison`, bold heading, caption and small credit."""
    cells, n = s['cells'], s.get('columns', 2)
    rows = math.ceil(len(cells) / n)
    gx, gy = 0.4, 0.25
    cw = (W - 2 * MX - gx * (n - 1)) / n
    ch = (BOTTOM - TOP - gy * (rows - 1)) / rows
    iw = cw * 0.42
    for i, c in enumerate(cells):
        r, k = divmod(i, n)
        x, y = MX + k * (cw + gx), TOP + r * (ch + gy)
        box(slide, x, y, cw, ch, LIGHT)
        if c.get('image'):
            pic = picture(slide, img(c['image']), x + 0.12, y + 0.12, h=ch - 0.24, name=f'yaml:cells.{i}.image')
            if pic.width.inches > iw:                          # crop around the centre to the image column
                keep = iw / pic.width.inches
                pic.crop_left = pic.crop_right = (1 - keep) / 2
                pic.width = Inches(iw)
        elif c.get('drawn') == 'comparison':
            draw_comparison(slide, x + 0.12, y + 0.12, iw, ch - 0.24, i)
        tx = x + iw + 0.32
        tw = cw - iw - 0.44
        para_box(slide, tx, y + 0.15, tw, 0.9, c['heading'], 20, INK, True, name=f'yaml:cells.{i}.heading')
        para_box(slide, tx, y + 1.0, tw, ch - 1.55, c['caption'], 18, INK, name=f'yaml:cells.{i}.caption')
        if c.get('credit'):
            para_box(slide, tx, y + ch - 0.5, tw, 0.4, c['credit'], 12, MUTED, name=f'yaml:cells.{i}.credit')


def draw_comparison(slide, x, y, w, h, i):
    """Illustrative response comparison: two response cards, the preferred one ticked (native shapes)."""
    bh = (h - 0.55) / 2
    for j, (label, mark) in enumerate((('Response A', '✓ preferred'), ('Response B', ''))):
        by = y + 0.45 + j * (bh + 0.1)
        b = box(slide, x, by, w, bh, WHITE, line=RGBColor(0xBF, 0xBF, 0xBF), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        b.name = f'deco:comparison.{i}.{j}'
        para_box(slide, x + 0.1, by + 0.05, w - 0.2, 0.35, label, 14, MUTED, True, name=f'deco:comparison.{i}.{j}.label')
        for k in range(2):
            ln = box(slide, x + 0.1, by + 0.45 + k * 0.22, w * (0.75 - 0.2 * k), 0.08, LIGHT)
            ln.name = f'deco:comparison.{i}.{j}.line{k}'
        if mark:
            para_box(slide, x + 0.1, by + bh - 0.45, w - 0.2, 0.4, mark, 14, RED, True,
                     name=f'deco:comparison.{i}.{j}.mark')
    para_box(slide, x, y, w, 0.4, 'Which response is better?', 14, INK, True, name=f'deco:comparison.{i}.prompt')


def k_native_bullets(slide, s, img):
    """Copy named, relationship-free shapes (e.g. a grouped diagram) from an old deck as native XML, centred at
    the top; optional `relabel` replaces whole text runs; bullets go below."""
    nat = s['native']
    src = Presentation(ROOT / nat['deck']).slides[nat['slide'] - 1]
    top = TOP
    for name in nat['shapes']:
        sh = next(x for x in src.shapes if x.name == name)
        el = copy.deepcopy(sh._element)
        if el.xpath('.//@r:embed | .//@r:id | .//@r:link'):
            raise ValueError(f"{s['id']}: native shape {name!r} has relationships; copy it through PowerPoint")
        used = {int(i) for i in slide.shapes._spTree.xpath('.//p:cNvPr/@id')}
        nxt = max(used | {1}) + 1
        for c in el.xpath('.//p:cNvPr'):                       # unique shape ids, or PowerPoint asks to repair
            c.set('id', str(nxt)); nxt += 1
        for t in el.xpath('.//a:t'):
            if t.text in nat.get('relabel', {}):
                t.text = nat['relabel'][t.text]
        slide.shapes._spTree.append(el)
        new = slide.shapes[-1]
        new.name = f"native:{nat['deck'].split('/')[-1]}#{nat['slide']}:{name}"
        new.left, new.top = Inches((W - new.width.inches) / 2), Inches(top)
        top += new.height.inches + 0.2
    bullets(slide, MX, top + 0.1, W - 2 * MX, BOTTOM - top - 0.1, s['bullets'], name='yaml:bullets',
            style=s.get('bullet_style', 'bullet'))


KINDS = {'title': k_title, 'course_title': k_course_title, 'image_bullets': k_image_bullets, 'bullets': k_bullets, 'two_column': k_two_column,
         'flow': k_flow, 'table': k_table, 'project': k_project, 'callouts': k_callouts, 'closing': k_closing,
         'section': k_closing, 'blank': k_blank, 'pptx_only': k_pptx_only, 'math': k_math, 'image_grid': k_image_grid, 'native_bullets': k_native_bullets}
LAYOUT = {'title': 'Title Slide', 'course_title': 'Title Slide', 'bullets': 'Title and Content', 'closing': 'Section Header',
          'section': 'Section Header'}  # else Title Only


# ---------------------------------------------------------------- build / render

def content_frame(prs):
    """Read the content frame (in) from the master's body placeholder, so the template is the one source."""
    global MX, TOP, BOTTOM
    body = next(ph for ph in prs.slide_master.placeholders if ph.placeholder_format.type == 2)
    MX, TOP, BOTTOM = body.left.inches, body.top.inches, Emu(body.top + body.height).inches


def make_presentation(spec):
    """Build the deck in memory (no validation or saving); shared with pull_deck_claude.py."""
    d = spec['deck']
    template = ROOT / d['template']
    src = Presentation(ROOT / d['image_source']) if d.get('image_source') else None

    def img(ref):
        """Image reference: {path: repo-relative file} or {source_slide: n, shape: name} in image_source."""
        if 'path' in ref:
            return (ROOT / ref['path']).read_bytes()
        shape = next(sh for sh in src.slides[ref['source_slide'] - 1].shapes if sh.name == ref['shape'])
        return shape.image.blob

    prs = Presentation(template)
    content_frame(prs)
    ids = prs.slides._sldIdLst
    for sid in list(ids):                       # drop template specimens; keep masters/layouts/theme
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    layouts = {l.name: l for l in prs.slide_layouts}
    for n, s in enumerate(spec['slides'], 1):
        slide = prs.slides.add_slide(layouts[LAYOUT.get(s['kind'], 'Title Only')])
        slide._idx = n
        slide._element.cSld.set('name', s['id'])           # spec id, for mapping PowerPoint edits back
        if s.get('hidden'):
            slide._element.set('show', '0')                # hidden in the slide show (e.g. a backup appendix)
        if s['kind'] == 'blank' and not s.get('title'):
            ph = slide.shapes.title                        # untitled live canvas: drop the empty title box
            ph._element.getparent().remove(ph._element)
        elif s['kind'] not in ('title', 'course_title', 'closing', 'section'):
            set_title(slide, s['title'])
        KINDS[s['kind']](slide, dict(s, _deck=d) if s['kind'] == 'course_title' else s, img)
        OMML.wrap_math_shapes(slide)
        notes(slide, s)
    cp = prs.core_properties
    cp.title, cp.author, cp.last_modified_by = d['title'], d['author'], 'scripts/build_deck_claude.py'
    cp.subject = f"Built from {d['template']} and {spec['_path']}"
    return prs


def build(spec, out):
    out = Path(out)
    validate_candidate_output(ROOT, out)
    prs = make_presentation(spec)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open('xb') as destination:
        prs.save(destination)


def preflight(pptx):
    """Refuse files PowerPoint might stop on with a dialog: every equation shape must sit in an
    mc:AlternateContent with a Choice and a Fallback, and shape ids must be unique per slide."""
    from lxml import etree
    problems = []
    for n, s in enumerate(Presentation(pptx).slides, 1):
        tree = s.shapes._spTree
        def in_fallback(el):
            while el is not None:
                if el.tag == OMML.MC + 'Fallback':
                    return True
                el = el.getparent()
            return False
        ids = [c.get('id') for c in tree.iter(qn('p:cNvPr')) if not in_fallback(c)]   # fallbacks repeat ids
        if len(ids) != len(set(ids)):
            problems.append(f'slide {n}: duplicate shape ids')
        for ac in tree.iter(OMML.MC + 'AlternateContent'):
            kids = [c.tag for c in ac]
            if kids != [OMML.MC + 'Choice', OMML.MC + 'Fallback']:
                problems.append(f'slide {n}: AlternateContent without Choice+Fallback')
        for m in tree.iter(OMML.A14 + 'm'):
            anc = m.getparent()
            while anc is not None and anc.tag != OMML.MC + 'Choice':
                anc = anc.getparent()
            if anc is None:
                problems.append(f'slide {n}: equation outside mc:Choice')
    if problems:
        raise SystemExit('preflight failed, not opening in PowerPoint:\n  ' + '\n  '.join(problems))


def render_dir(name):
    """PowerPoint (sandboxed) opens files in the system temp folder without asking; other folders, such as the
    agent's scratch space, trigger a blocking 'grant access' dialog. Always render here."""
    d = Path(tempfile.gettempdir()) / 'deck_render' / name
    d.mkdir(parents=True, exist_ok=True)
    return d


def render(pptx, preview, workdir):
    preflight(pptx)
    if not Path(workdir).resolve().is_relative_to(Path(tempfile.gettempdir()).resolve()):
        raise SystemExit(f'render folder {workdir} is outside the system temp folder; PowerPoint would ask for '
                         'access and block. Use the default (no --render-dir).')
    script = ('on run argv\n tell application "Microsoft PowerPoint"\n'
              '  open (POSIX file (item 1 of argv))\n'
              '  set p to presentation (item 3 of argv)\n'
              '  save p in (POSIX file (item 2 of argv)) as save as PDF\n  close p saving no\n'
              ' end tell\nend run\n')
    work = Path(workdir)
    # Reuse one folder per deck: remove only this builder's own previous outputs, so disk use stays
    # bounded and stale slide images from a longer earlier build cannot leak into the contact sheet.
    for old in [*work.glob('slide-*.png'), *work.glob('render_*.pptx'), *work.glob('render_*.pdf')]:
        old.unlink(missing_ok=True)
    import os, time
    stem = f'render_{os.getpid()}_{int(time.time())}'   # unique: never picks up a presentation left open earlier
    copy_ = work / f'{stem}.pptx'              # never open the repository file in PowerPoint
    copy_.write_bytes(Path(pptx).read_bytes())
    pdf = work / f'{stem}.pdf'
    subprocess.run(['osascript', '-e', script, str(copy_), str(pdf), copy_.name], check=True, timeout=180)
    subprocess.run(['pdftoppm', '-r', '110', '-png', str(pdf), str(work / 'slide')], check=True)
    from PIL import Image
    pages = sorted(work.glob('slide-*.png'))
    ims = [Image.open(p).convert('RGB') for p in pages]
    tw = 640
    thumbs = [im.resize((tw, round(im.height * tw / im.width))) for im in ims]
    cols, pad = 3, 10
    th = thumbs[0].height
    rows = math.ceil(len(thumbs) / cols)
    sheet = Image.new('RGB', (cols * (tw + pad) + pad, rows * (th + pad) + pad), (110, 110, 110))
    for i, t in enumerate(thumbs):
        sheet.paste(t, (pad + (i % cols) * (tw + pad), pad + (i // cols) * (th + pad)))
    Path(preview).parent.mkdir(parents=True, exist_ok=True)
    sheet.save(preview, optimize=True)
    return pages


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('spec')
    ap.add_argument('--check', action='store_true', help='validate the spec only; write nothing')
    ap.add_argument('--render', action='store_true', help='also render via PowerPoint and write the preview')
    ap.add_argument('--render-dir', help='directory for render intermediates '
                    '(default: <system temp>/deck_render/<spec name>, reused and overwritten each render)')
    a = ap.parse_args()
    spec_path = Path(a.spec).resolve()
    spec = yaml.safe_load(spec_path.read_text(encoding='utf-8'))
    spec['_path'] = str(spec_path.relative_to(ROOT))
    errors, used, budget = validate(spec)
    nq = sum(len(s.get('verify', [])) for s in spec['slides'])
    print(f'{len(spec["slides"])} slides; {nq} syllabus quotes checked; timing {used} + '
          f'{spec["deck"].get("qa_minutes", 0)} Q&A of {budget} min')
    if errors:
        print('\n'.join('ERROR ' + e for e in errors))
        sys.exit(1)
    if a.check:
        for w_ in warnings:
            print('WARN', w_)
        return
    out = candidate_output(ROOT, spec['deck']['output'])
    build(spec, out)
    print(f'wrote {out.relative_to(ROOT)}')
    for w_ in warnings:
        print('WARN', w_)
    if a.render:
        work = a.render_dir or Path(tempfile.gettempdir()) / 'deck_render' / spec_path.stem
        Path(work).mkdir(parents=True, exist_ok=True)
        pages = render(out, ROOT / spec['deck']['preview'], work)
        (Path(work) / 'last_render.pdf').write_bytes(sorted(Path(work).glob('render_*.pdf'))[-1].read_bytes())
        print(f'rendered {len(pages)} pages to {work}; preview {spec["deck"]["preview"]}')


if __name__ == '__main__':
    main()
