#!/usr/bin/env python3
"""Apply the course's equation style to a saved lecture deck (Claude-built workflow).

    python3 scripts/deck_math_style_claude.py <deck.pptx> [--dry-run] [--no-color] [--no-superscripts]

1. Colour: every native equation run (OMML) uses the theme colour Accent 5 (instructor, 2026-10-11; #4472C4 in the
   course template), replacing any other run colour inside equations.
2. Superscripts (notation guide, instructor 2026-10-11): policy and optimality marks move from subscripts to
   superscripts: v_π → v^π, q_π → q^π, v_* → v^*, q_* → q^*, π_* → π^*, B_π → B^π. Applied to native equations, to
   math typeset as text runs (tables, diagram labels) and to the speaker notes (plain text and LaTeX sources).
   Other subscripts (𝔼_π, V_k, S_t, π_θ, ...) are unchanged.

The deck must be closed in PowerPoint. A backup goes to output/deck_backups/ first. Equations edited here keep stale
fallback pictures; refresh them with scripts/refresh_math_fallbacks_claude.py before the final check.
"""
import argparse, hashlib, re, shutil, sys, time
from pathlib import Path
from lxml import etree
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
BASES = {'v', 'q', 'B', 'π', '𝑣', '𝑞', '𝐵', '𝜋'}
MARKS = {'π', '𝜋', '*', '∗'}

# notes: plain text and LaTeX sources
NOTE_SUBS = [
    (re.compile(r'(?<![A-Za-z\\])([vqB])_\\pi\b'), r'\1^{\\pi}'),
    (re.compile(r'(?<![A-Za-z\\])([vq])_\*'), r'\1^{*}'),
    (re.compile(r'\\pi_\*'), r'\\pi^{*}'),
    (re.compile(r'(?<![A-Za-z])([vqB])_π'), r'\1^π'),
    (re.compile(r'(?<![A-Za-z])([vq])_\*'), r'\1^*'),
    (re.compile(r'π_\*'), r'π^*'),
    (re.compile(r'(?<![A-Za-z])([vqB])π(?=[\s(,.;:)]|$)'), r'\1^π'),   # flattened "vπ(s)" written in notes
]


def text_of(el):
    return ''.join(el.itertext()).strip()


def recolor(root):
    n = 0
    for om in root.iter(M + 'oMath'):
        for rpr in om.iter(A + 'rPr'):
            for f in rpr.findall(A + 'solidFill'):
                rpr.remove(f)
            fill = etree.Element(A + 'solidFill')
            etree.SubElement(fill, A + 'schemeClr', val='accent5')
            # schema order: ln, fill, effect..., latin/ea/cs: put the fill before any font element
            fonts = [c for c in rpr if c.tag in (A + 'latin', A + 'ea', A + 'cs', A + 'sym', A + 'ln')]
            if fonts and fonts[0].tag != A + 'ln':
                fonts[0].addprevious(fill)
            elif fonts:
                fonts[0].addnext(fill)
            else:
                rpr.append(fill)
            n += 1
    return n


def omml_superscripts(root, log):
    n = 0
    for s in list(root.iter(M + 'sSub')):
        e, sub = s.find(M + 'e'), s.find(M + 'sub')
        if e is None or sub is None or text_of(e) not in BASES or text_of(sub) not in MARKS:
            continue
        s.tag = M + 'sSup'
        pr = s.find(M + 'sSubPr')
        if pr is not None:
            pr.tag = M + 'sSupPr'
        sub.tag = M + 'sup'
        log.append(f'{text_of(e)}_{text_of(sub)} -> ^'); n += 1
    for s in root.iter(M + 'sSubSup'):
        e, sub = s.find(M + 'e'), s.find(M + 'sub')
        if e is not None and sub is not None and text_of(e) in BASES and text_of(sub) in MARKS:
            log.append(f'UNCHANGED sSubSup {text_of(e)}_{text_of(sub)}^{text_of(s.find(M + "sup"))}')
    return n


def textrun_superscripts(root, log):
    """Math typeset as text runs (builder's text_math): a lowered π/* run right after a v/q/B run is raised."""
    n = 0
    for p in root.iter(A + 'p'):
        runs = p.findall(A + 'r')
        for prev, r in zip(runs, runs[1:]):
            rpr, t, pt = r.find(A + 'rPr'), r.findtext(A + 't') or '', prev.findtext(A + 't') or ''
            if rpr is not None and t.strip() in MARKS and int(rpr.get('baseline', '0')) < 0 \
                    and pt.strip()[-1:] in BASES:
                rpr.set('baseline', '30000'); log.append(f'text run {pt.strip()[-1]}_{t.strip()} -> ^'); n += 1
    return n


def notes_superscripts(slide, log):
    if not slide.has_notes_slide:
        return 0
    n = 0
    for t in slide.notes_slide.notes_text_frame._txBody.iter(A + 't'):
        new = t.text or ''
        for pat, rep in NOTE_SUBS:
            new = pat.sub(rep, new)
        if new != t.text:
            log.append(f'notes: {t.text[:60]!r} -> {new[:60]!r}'); t.text = new; n += 1
    return n


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('deck')
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--no-color', action='store_true')
    ap.add_argument('--no-superscripts', action='store_true')
    a = ap.parse_args()
    deck = Path(a.deck)
    lock = deck.with_name('~$' + deck.name)
    if lock.exists():
        sys.exit(f'{deck.name} is open in PowerPoint ({lock.name}); save and close it first.')
    prs = Presentation(deck)
    total = {'color': 0, 'omml': 0, 'text': 0, 'notes': 0}
    for i, s in enumerate(prs.slides, 1):
        log = []
        root = s._element
        if not a.no_color:
            total['color'] += recolor(root)
        if not a.no_superscripts:
            total['omml'] += omml_superscripts(root, log)
            total['text'] += textrun_superscripts(root, log)
            total['notes'] += notes_superscripts(s, log)
        for line in log:
            print(f'slide {i} ({s._element.cSld.get("name")}): {line}')
    print('changed:', total)
    if a.dry_run:
        return
    h = hashlib.sha256(deck.read_bytes()).hexdigest()[:8]
    bak = ROOT / 'output/deck_backups' / f'{deck.stem}.{time.strftime("%Y%m%d-%H%M%S")}.{h}.pptx'
    shutil.copy2(deck, bak)
    prs.save(deck)
    print(f'backup {bak.relative_to(ROOT)}; saved {deck}')


if __name__ == '__main__':
    main()
