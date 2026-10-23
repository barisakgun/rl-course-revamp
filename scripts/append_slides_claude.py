#!/usr/bin/env python3
"""Append the slides of a built candidate deck to an existing (instructor-edited) deck (Claude-built workflow).

    python3 scripts/append_slides_claude.py <deck.pptx> <candidate.pptx> [--expect-sha SHA] [--dry-run]

Used when a later session of a multi-session deck is rendered: the new slides are built from their own YAML
(build_deck_claude.py), their equation fallbacks refreshed on the candidate, and then appended here. The existing
slides are never rebuilt: only new slide parts are added, so the existing slides' XML is preserved.

For each candidate slide: the slide layout with the same name in the deck is used; the shape tree, slide name, slide
attributes, builds/animations (p:timing) and transition are copied; images (including equation fallback pictures)
are added to the deck and their relationship ids remapped; the speaker notes body is copied.

--expect-sha refuses to run if the deck is not the expected saved version (e.g. the instructor saved new edits).
A backup goes to output/deck_backups/ before the deck is replaced. The deck must not be open in PowerPoint.
"""
import argparse, copy, hashlib, io, shutil, sys, time
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
RT_IMAGE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def append(deck, cand):
    layouts = {l.name: l for l in deck.slide_layouts}
    for cs in cand.slides:
        lay = layouts[cs.slide_layout.name]
        ns = deck.slides.add_slide(lay)
        tree = ns.shapes._spTree
        for el in list(tree)[2:]:                          # drop the layout placeholders the new slide received
            tree.remove(el)
        # images: add each to the new slide part, remember the id mapping
        remap = {}
        for rid, rel in cs.part.rels.items():
            if rel.reltype == RT_IMAGE:
                _, new = ns.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))
                remap[rid] = new
            elif rel.reltype.split('/')[-1] not in ('slideLayout', 'notesSlide'):
                raise SystemExit(f'unsupported relationship {rel.reltype} on candidate slide {cs.slide_id}')
        for el in list(cs.shapes._spTree)[2:]:
            el = copy.deepcopy(el)
            for node in el.iter():
                for attr in list(node.attrib):
                    if attr.startswith(f'{{{R_NS}}}') and node.get(attr) in remap:
                        node.set(attr, remap[node.get(attr)])
            tree.append(el)
        # slide name (stable id), attributes (e.g. show="0"), timing and transition
        src, dst = cs._element, ns._element
        if src.cSld.get('name'):
            dst.cSld.set('name', src.cSld.get('name'))
        for k, v in src.attrib.items():
            dst.set(k, v)
        anchor = dst.find(qn('p:clrMapOvr'))
        for tag in ('p:transition', 'p:timing'):
            el = src.find(qn(tag))
            if el is not None:
                el = copy.deepcopy(el)
                (anchor if anchor is not None else dst.cSld).addnext(el)
                anchor = el
        # speaker notes
        if cs.has_notes_slide:
            body = cs.notes_slide.notes_text_frame._txBody
            target = ns.notes_slide.notes_text_frame._txBody
            for el in list(target):
                target.remove(el)
            for el in body:
                target.append(copy.deepcopy(el))
    return len(cand.slides)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('deck')
    ap.add_argument('candidate')
    ap.add_argument('--expect-sha')
    ap.add_argument('--dry-run', action='store_true', help='write <deck>.appended.pptx in the system temp folder only')
    a = ap.parse_args()
    deck_path = Path(a.deck)
    lock = deck_path.with_name('~$' + deck_path.name)
    if lock.exists():
        raise SystemExit(f'{lock.name} exists: the deck appears to be open in PowerPoint. Close it first.')
    before = sha(deck_path)
    if a.expect_sha and before != a.expect_sha:
        raise SystemExit(f'deck SHA-256 {before} differs from the expected {a.expect_sha}; reconcile first.')
    deck, cand = Presentation(deck_path), Presentation(a.candidate)
    n0 = len(deck.slides)
    n = append(deck, cand)
    buf = io.BytesIO()
    deck.save(buf)
    if a.dry_run:
        import tempfile
        out = Path(tempfile.gettempdir()) / f'{deck_path.stem}.appended.pptx'
        out.write_bytes(buf.getvalue())
        print(f'dry run: {n0} + {n} slides -> {out}')
        return
    backups = ROOT / 'output' / 'deck_backups'
    backups.mkdir(parents=True, exist_ok=True)
    bk = backups / f'{deck_path.stem}.{time.strftime("%Y%m%d-%H%M%S")}.{before[:8]}.pptx'
    shutil.copy2(deck_path, bk)
    deck_path.write_bytes(buf.getvalue())
    print(f'appended {n} slides after slide {n0}; backup {bk.relative_to(ROOT)}; new SHA-256 {sha(deck_path)}')


if __name__ == '__main__':
    main()
