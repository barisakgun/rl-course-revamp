#!/usr/bin/env python3
"""Build the course lecture template (slide master + layouts) used by build_deck_claude.py.

    python3 scripts/build_template_claude.py [--render]

Starts from the Office theme of the instructor's existing decks (so pasted Spring 2025 slides map onto
it cleanly), then rewrites the slide master and layouts:

  * Title: top-left corner at (LEFT, TITLE_TOP), height TITLE_H, Calibri Light 36 pt, dark red, no insets.
  * A full-width red rule just below the title; the body starts GAP_RULE_BODY below the rule's bottom edge.
  * Body: left edge aligned with the title, bullet levels styled on the master (• red, – grey, • grey).
  * Footer text and an auto-updating slide-number field are master shapes, so pasted slides get them too.
  * Layout names match the Spring 2025 decks ("Title and Content", "Title Slide", "Title, Bullets & Photo"),
    which PowerPoint uses to map pasted slides with "Use Destination Theme".
  * Theme colours and fonts stay Office defaults so diagrams pasted from old decks keep their colours.

The template contains one specimen slide per layout for visual checks; the deck builder drops them.
Output: course/templates/lecture_template_claude.pptx (refuses to overwrite; delete it explicitly first).
"""
import argparse, copy, subprocess, sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'sources/current_course/0 - NutsAndBolts.pptx'    # Office theme of the existing decks
OUT = ROOT / 'course/templates/lecture_template_claude.pptx'

# ---- geometry (mm); everything else is derived -------------------------------------------------
SLIDE_W, SLIDE_H = 338.667, 190.5
LEFT = 6.2               # title and body left edge; also used as the right margin
TITLE_TOP = 3.5          # top of the title box
TITLE_H = 12.7           # ≈ one 36 pt line (36 pt = 12.7 mm)
GAP_TITLE_RULE = 1.0     # title box bottom -> rule top
RULE_H = 0.75            # rule thickness (≈ 2.1 pt)
GAP_RULE_BODY = 3.5      # rule bottom -> body top
FOOTER_H = 4.5           # footer strip height
FOOTER_BOTTOM = 2.0      # footer strip bottom -> slide bottom
GAP_BODY_FOOTER = 1.5    # body bottom -> footer strip top
COL_GAP = 8.0            # gap between two content columns
FOOTER_TEXT = 'COMP438/538 · Fall 2026'

RULE_TOP = TITLE_TOP + TITLE_H + GAP_TITLE_RULE
BODY_TOP = RULE_TOP + RULE_H + GAP_RULE_BODY
FOOTER_TOP = SLIDE_H - FOOTER_BOTTOM - FOOTER_H
BODY_BOTTOM = FOOTER_TOP - GAP_BODY_FOOTER
CONTENT_W = SLIDE_W - 2 * LEFT
BODY_H = BODY_BOTTOM - BODY_TOP
COL_W = (CONTENT_W - COL_GAP) / 2

RED, GREY, MUTED = 'C00000', '7F7F7F', '595959'
# (size pt, bullet char, bullet colour, left margin mm, space before pt)
LEVELS = [(24, '•', RED, 6.35, 6), (20, '–', GREY, 14.3, 3), (18, '•', GREY, 21.6, 2),
          (16, '–', GREY, 28.6, 2), (16, '–', GREY, 35.6, 2)]
HANG = 6.35              # bullet hang (mm): bullet at the box edge, text at marL

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def emu(mm):
    return str(int(round(mm * 36000)))


def el(tag, attrib=None, *children):
    e = etree.Element(qn(tag), {k: str(v) for k, v in (attrib or {}).items()})
    for c in children:
        e.append(c)
    return e


def solid(hex_):
    return el('a:solidFill', None, el('a:srgbClr', {'val': hex_}))


def set_xfrm(sp, x, y, w, h):
    spPr = sp.find(qn('p:spPr'))
    for old in spPr.findall(qn('a:xfrm')):
        spPr.remove(old)
    spPr.insert(0, el('a:xfrm', None, el('a:off', {'x': emu(x), 'y': emu(y)}),
                      el('a:ext', {'cx': emu(w), 'cy': emu(h)})))


def drop_xfrm(sp):
    spPr = sp.find(qn('p:spPr'))
    for old in spPr.findall(qn('a:xfrm')):
        spPr.remove(old)


def set_body_pr(sp, anchor=None, autofit='norm'):
    txBody = sp.find(qn('p:txBody'))
    bp = txBody.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'):
        bp.set(k, '0')
    if anchor:
        bp.set('anchor', anchor)
    for c in list(bp):
        if c.tag in (qn('a:normAutofit'), qn('a:noAutofit'), qn('a:spAutoFit')):
            bp.remove(c)
    if autofit == 'norm':
        bp.append(el('a:normAutofit'))


def strip_lst_style(sp):
    lst = sp.find(qn('p:txBody')).find(qn('a:lstStyle'))
    if lst is not None:
        for c in list(lst):
            lst.remove(c)


def lst_style(sp, xml_levels):
    lst = sp.find(qn('p:txBody')).find(qn('a:lstStyle'))
    for c in list(lst):
        lst.remove(c)
    for lv in xml_levels:
        lst.append(lv)


def ph_type(sp):
    ph = sp.find('.//' + qn('p:ph'))
    if ph is None:
        return None
    return ph.get('type', 'obj'), ph.get('idx', '0')


def remove_footer_placeholders(tree):
    for sp in tree.findall('.//' + qn('p:sp')):
        t = ph_type(sp)
        if t and t[0] in ('dt', 'ftr', 'sldNum'):
            sp.getparent().remove(sp)


def title_level():
    rpr = el('a:defRPr', {'sz': 3600, 'kern': 1200}, solid(RED), el('a:latin', {'typeface': '+mj-lt'}),
             el('a:ea', {'typeface': '+mj-ea'}), el('a:cs', {'typeface': '+mj-cs'}))
    return el('a:lvl1pPr', {'algn': 'l', 'defTabSz': 914400, 'rtl': 0, 'eaLnBrk': 1, 'latinLnBrk': 0,
                            'hangingPunct': 1},
              el('a:lnSpc', None, el('a:spcPct', {'val': 90000})),
              el('a:spcBef', None, el('a:spcPct', {'val': 0})), el('a:buNone'), rpr)


def body_level(i):
    size, char, colour, marl, before = LEVELS[min(i, len(LEVELS) - 1)]
    rpr = el('a:defRPr', {'sz': size * 100, 'kern': 1200}, el('a:solidFill', None, el('a:schemeClr', {'val': 'tx1'})),
             el('a:latin', {'typeface': '+mn-lt'}), el('a:ea', {'typeface': '+mn-ea'}),
             el('a:cs', {'typeface': '+mn-cs'}))
    return el(f'a:lvl{i + 1}pPr', {'marL': emu(marl), 'indent': emu(-HANG), 'algn': 'l', 'defTabSz': 914400,
                                  'rtl': 0, 'eaLnBrk': 1, 'latinLnBrk': 0, 'hangingPunct': 1},
              el('a:lnSpc', None, el('a:spcPct', {'val': 100000})),
              el('a:spcBef', None, el('a:spcPts', {'val': before * 100})),
              el('a:buClr', None, el('a:srgbClr', {'val': colour})),
              el('a:buFont', {'typeface': 'Arial', 'panose': '020B0604020202020204', 'pitchFamily': 34,
                              'charset': 0}),
              el('a:buChar', {'char': char}), rpr)


def plain_level(i, size, colour=None, algn='l'):
    kids = [el('a:solidFill', None, el('a:srgbClr', {'val': colour}))] if colour else []
    return el(f'a:lvl{i + 1}pPr', {'marL': 0, 'indent': 0, 'algn': algn}, el('a:buNone'),
              el('a:defRPr', {'sz': size * 100}, *kids))


def text_shape(spid, name, x, y, w, h, size, colour, field=None, text='', algn='l'):
    """Non-placeholder text box (master decoration) with optional slide-number field."""
    r_pr = el('a:rPr', {'lang': 'en-US', 'sz': size * 100, 'dirty': 0}, solid(colour))
    if field:
        run = el('a:fld', {'id': '{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}', 'type': field}, r_pr,
                 el('a:t'))
        run.find(qn('a:t')).text = '‹#›'
    else:
        run = el('a:r', None, r_pr, el('a:t'))
        run.find(qn('a:t')).text = text
    para = el('a:p', None, el('a:pPr', {'algn': algn}), run)
    sp = el('p:sp', None,
            el('p:nvSpPr', None, el('p:cNvPr', {'id': spid, 'name': name}), el('p:cNvSpPr', {'txBox': 1}),
               el('p:nvPr')),
            el('p:spPr', None, el('a:xfrm', None, el('a:off', {'x': emu(x), 'y': emu(y)}),
                                  el('a:ext', {'cx': emu(w), 'cy': emu(h)})),
               el('a:prstGeom', {'prst': 'rect'}, el('a:avLst')), el('a:noFill')),
            el('p:txBody', None, el('a:bodyPr', {'wrap': 'square', 'lIns': 0, 'tIns': 0, 'rIns': 0, 'bIns': 0,
                                                 'anchor': 'ctr'}, el('a:noAutofit')), el('a:lstStyle'), para))
    return sp


def rule_shape(spid):
    return el('p:sp', None,
              el('p:nvSpPr', None, el('p:cNvPr', {'id': spid, 'name': 'Title rule'}), el('p:cNvSpPr'), el('p:nvPr')),
              el('p:spPr', None, el('a:xfrm', None, el('a:off', {'x': 0, 'y': emu(RULE_TOP)}),
                                    el('a:ext', {'cx': emu(SLIDE_W), 'cy': emu(RULE_H)})),
                 el('a:prstGeom', {'prst': 'rect'}, el('a:avLst')), solid(RED), el('a:ln', None, el('a:noFill'))))


def build_master(master):
    m = master._element
    tree = m.find(qn('p:cSld')).find(qn('p:spTree'))
    remove_footer_placeholders(tree)
    for sp in tree.findall(qn('p:sp')):
        t = ph_type(sp)
        if t and t[0] == 'title':
            set_xfrm(sp, LEFT, TITLE_TOP, CONTENT_W, TITLE_H)
            set_body_pr(sp, anchor='b')
        elif t and t[0] == 'body':
            set_xfrm(sp, LEFT, BODY_TOP, CONTENT_W, BODY_H)
            set_body_pr(sp, anchor='t')
    ids = [int(e.get('id')) for e in m.iter(qn('p:cNvPr'))]
    nid = max(ids) + 1
    tree.append(rule_shape(nid))
    tree.append(text_shape(nid + 1, 'Footer text', LEFT, FOOTER_TOP, 120, FOOTER_H, 10, MUTED, text=FOOTER_TEXT))
    tree.append(text_shape(nid + 2, 'Slide number', SLIDE_W - LEFT - 30, FOOTER_TOP, 30, FOOTER_H, 10, MUTED,
                           field='slidenum', algn='r'))
    styles = m.find(qn('p:txStyles'))
    ts = styles.find(qn('p:titleStyle'))
    for c in list(ts):
        ts.remove(c)
    ts.append(title_level())
    bs = styles.find(qn('p:bodyStyle'))
    for c in list(bs):
        bs.remove(c)
    for i in range(9):
        bs.append(body_level(i))


def layout_placeholders(layout):
    out = {}
    for sp in layout._element.find(qn('p:cSld')).find(qn('p:spTree')).findall(qn('p:sp')):
        t = ph_type(sp)
        if t:
            out.setdefault(t[0], []).append(sp)
    return out


def hide_master_shapes(layout):
    layout._element.set('showMasterSp', '0')


def build_layouts(prs):
    master = prs.slide_master
    keep = {'Title Slide', 'Title and Content', 'Section Header', 'Two Content', 'Comparison', 'Title Only',
            'Blank', 'Picture with Caption'}
    for layout in list(master.slide_layouts):
        if layout.name not in keep:
            master.slide_layouts.remove(layout)
    for layout in master.slide_layouts:
        remove_footer_placeholders(layout._element)
        ph = layout_placeholders(layout)
        name = layout.name
        for sp in ph.get('title', []) + ph.get('ctrTitle', []):
            if name not in ('Title Slide', 'Section Header', 'Picture with Caption'):
                drop_xfrm(sp)                      # inherit master title geometry and style
                strip_lst_style(sp)
        if name in ('Title and Content', 'Title Only'):
            for sp in ph.get('obj', []) + ph.get('body', []):
                drop_xfrm(sp)
                strip_lst_style(sp)
        elif name == 'Two Content':
            for i, sp in enumerate(ph.get('obj', [])):
                set_xfrm(sp, LEFT + i * (COL_W + COL_GAP), BODY_TOP, COL_W, BODY_H)
                strip_lst_style(sp)
                set_body_pr(sp, anchor='t')
        elif name == 'Comparison':
            heads, bodies = ph.get('body', []), ph.get('obj', [])
            for i, sp in enumerate(heads):
                set_xfrm(sp, LEFT + i * (COL_W + COL_GAP), BODY_TOP, COL_W, 11)
                set_body_pr(sp, anchor='b', autofit=None)
                lst_style(sp, [el('a:lvl1pPr', {'marL': 0, 'indent': 0}, el('a:buNone'),
                                  el('a:defRPr', {'sz': 2400, 'b': 1}))])
            for i, sp in enumerate(bodies):
                set_xfrm(sp, LEFT + i * (COL_W + COL_GAP), BODY_TOP + 13, COL_W, BODY_H - 13)
                strip_lst_style(sp)
                set_body_pr(sp, anchor='t')
        elif name == 'Picture with Caption':
            layout._element.cSld.set('name', 'Title, Bullets & Photo')   # name used by the Spring 2025 MDP deck
            for sp in ph.get('title', []):
                drop_xfrm(sp)
                strip_lst_style(sp)
                bp = sp.find(qn('p:txBody')).find(qn('a:bodyPr'))
                for k in list(bp.attrib):
                    bp.attrib.pop(k)
            for sp in ph.get('body', []):
                set_xfrm(sp, LEFT, BODY_TOP, COL_W, BODY_H)
                strip_lst_style(sp)
                set_body_pr(sp, anchor='t')
            for sp in ph.get('pic', []):
                set_xfrm(sp, LEFT + COL_W + COL_GAP, BODY_TOP, COL_W, BODY_H)
        elif name == 'Title Slide':
            hide_master_shapes(layout)
            for sp in ph.get('ctrTitle', []):
                set_xfrm(sp, LEFT, 52, CONTENT_W, 40)
                set_body_pr(sp, anchor='b')
                lst_style(sp, [plain_level(0, 44, RED)])
            for sp in ph.get('subTitle', []):
                set_xfrm(sp, LEFT, 96, CONTENT_W, 30)
                set_body_pr(sp, anchor='t')
                lst_style(sp, [plain_level(i, 24, MUTED) for i in range(1)])
        elif name == 'Section Header':
            hide_master_shapes(layout)
            for sp in ph.get('title', []):
                set_xfrm(sp, LEFT, 60, CONTENT_W, 40)
                set_body_pr(sp, anchor='b')
                lst_style(sp, [plain_level(0, 48, RED)])
            for sp in ph.get('body', []):
                set_xfrm(sp, LEFT, 104, CONTENT_W, 25)
                set_body_pr(sp, anchor='t')
                lst_style(sp, [plain_level(0, 24, MUTED)])
        elif name == 'Blank':
            hide_master_shapes(layout)


SPECIMEN = {
    'Title Slide': ('Lecture title', ['Course · term · instructor']),
    'Title and Content': ('Title and Content — default content layout',
                          ['Level 1: 24 pt, red bullet at the title’s left edge',
                           ['Level 2: 20 pt, grey dash', ['Level 3: 18 pt']],
                           f'Title box: {LEFT} mm from left, {TITLE_TOP} mm from top, {TITLE_H} mm high',
                           f'Rule: {RULE_TOP:.2f}–{RULE_TOP + RULE_H:.2f} mm; body starts at {BODY_TOP:.2f} mm',
                           f'Body: {CONTENT_W:.1f} × {BODY_H:.1f} mm']),
    'Two Content': ('Two Content', ['Left column'], ['Right column']),
    'Comparison': ('Comparison', ['Left heading'], ['Left content'], ['Right heading'], ['Right content']),
    'Title, Bullets & Photo': ('Title, Bullets & Photo', ['Bullets on the left; picture placeholder on the right']),
    'Title Only': ('Title Only — free layout area below the rule', []),
    'Section Header': ('Section header', ['Short description']),
    'Blank': None,
}


def fill(tf, items, level=0, first=None):
    first = first if first is not None else [True]
    for it in items:
        if isinstance(it, list):
            fill(tf, it, level + 1, first)
            continue
        p = tf.paragraphs[0] if first[0] else tf.add_paragraph()
        first[0] = False
        p.text, p.level = it, level


def add_specimens(prs):
    for layout in prs.slide_master.slide_layouts:
        spec = SPECIMEN.get(layout.name)
        slide = prs.slides.add_slide(layout)
        if spec is None:
            continue
        title, *bodies = spec
        if slide.shapes.title is not None:
            slide.shapes.title.text = title
        others = [p for p in slide.placeholders if p.placeholder_format.type in (2, 4, 7)]  # body/subtitle/object
        others.sort(key=lambda p: p.placeholder_format.idx)
        for ph, items in zip(others, bodies):
            if ph.has_text_frame:
                fill(ph.text_frame, items)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--render', action='store_true', help='render specimens via build_deck_claude.render')
    a = ap.parse_args()
    if OUT.exists():
        sys.exit(f'{OUT.relative_to(ROOT)} exists; delete it explicitly before rebuilding.')
    prs = Presentation(SOURCE)
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    build_master(prs.slide_master)
    build_layouts(prs)
    theme_part = prs.slide_master.part.part_related_by(
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme')
    theme = etree.fromstring(theme_part.blob)
    theme.set('name', 'COMP438 lecture (Claude)')
    theme_part._blob = etree.tostring(theme, xml_declaration=True, encoding='UTF-8', standalone=True)
    add_specimens(prs)
    cp = prs.core_properties
    cp.title, cp.last_modified_by = 'COMP438/538 lecture template', 'scripts/build_template_claude.py'
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open('xb') as f:
        prs.save(f)
    print(f'wrote {OUT.relative_to(ROOT)}: title top-left ({LEFT}, {TITLE_TOP}) mm, rule {RULE_TOP:.2f}+{RULE_H} mm, '
          f'body {BODY_TOP:.2f}–{BODY_BOTTOM:.2f} mm ({CONTENT_W:.1f} × {BODY_H:.1f} mm)')
    if a.render:
        sys.path.insert(0, str(Path(__file__).parent))
        from build_deck_claude import render
        import tempfile
        work = Path(tempfile.gettempdir()) / 'deck_render' / OUT.stem
        work.mkdir(parents=True, exist_ok=True)
        pages = render(OUT, ROOT / 'output/lectures/templates/lecture_template_claude_preview.png', work)
        print(f'rendered {len(pages)} specimen slides to {work}')


if __name__ == '__main__':
    main()
